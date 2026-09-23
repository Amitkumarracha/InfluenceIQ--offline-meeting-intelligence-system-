import os
from pptx import Presentation

def generate_ppt():
    prs = Presentation()
    
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "New Product Design"
    slide.placeholders[1].text = "Initial Thoughts & Remote Control"
    
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "The Interface"
    slide.placeholders[1].text_frame.text = "Should it be a standard remote control?"
    
    os.makedirs("data/custom_test", exist_ok=True)
    prs.save("data/custom_test/unseen_slides.pptx")
    print("Created test presentation: data/custom_test/unseen_slides.pptx")

if __name__ == "__main__":
    generate_ppt()
