import os
from pptx import Presentation
from pptx.util import Inches
from gtts import gTTS
from pydub import AudioSegment

def generate_ppt():
    prs = Presentation()
    
    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Project Phoenix Launch"
    subtitle.text = "Q3 Strategy Meeting"
    
    # Slide 2: Problem
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "The Current Problem"
    tf = body_shape.text_frame
    tf.text = "Our current systems are too slow."
    
    # Slide 3: Solution
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "The Solution: Phoenix"
    tf = body_shape.text_frame
    tf.text = "Phoenix will increase speed by 40%."
    
    os.makedirs("data/custom_test", exist_ok=True)
    prs.save("data/custom_test/phoenix_pitch.pptx")
    print("Created test presentation: data/custom_test/phoenix_pitch.pptx")

def generate_audio():
    # We will generate separate audio files for two "speakers" and concatenate them with pauses
    script = [
        ("en-us", "Welcome everyone to the Q3 Strategy Meeting. Today we are discussing Project Phoenix.", 0),
        ("en-co", "Thanks for having me. I see on the first slide that we are focusing on system speed?", 1000),
        ("en-us", "Exactly. As you can see on the next slide, the main problem is that our current systems are way too slow.", 500),
        ("en-co", "Right, I agree. So what is the proposed solution?", 1000),
        ("en-us", "Moving to the final slide, The Solution is Phoenix. It will increase our speed by forty percent. Do we have approval to move forward?", 500),
        ("en-co", "Yes, I approve the timeline for Project Phoenix.", 1000)
    ]
    
    combined = AudioSegment.empty()
    
    for i, (lang, text, pause) in enumerate(script):
        tts = gTTS(text, tld=lang.split('-')[1] if '-' in lang else 'com', lang='en')
        tmp_file = f"tmp_{i}.mp3"
        tts.save(tmp_file)
        segment = AudioSegment.from_mp3(tmp_file)
        combined += segment
        if pause > 0:
            combined += AudioSegment.silent(duration=pause)
        os.remove(tmp_file)
    
    # Export as wav
    combined = combined.set_frame_rate(16000).set_channels(1)
    combined.export("data/custom_test/phoenix_meeting.wav", format="wav")
    print("Created test audio: data/custom_test/phoenix_meeting.wav")

if __name__ == "__main__":
    generate_ppt()
    try:
        generate_audio()
    except Exception as e:
        print(f"Audio generation failed (maybe missing ffmpeg/pydub): {e}")
