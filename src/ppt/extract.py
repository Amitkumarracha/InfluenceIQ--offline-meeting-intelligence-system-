"""
PPT/PPTX slide extraction.

Provides:
- SlideData      : typed dataclass for a single slide
- PresentationData : typed dataclass for the full presentation
- PPTXExtractor  : adapter for standard .pptx files (uses python-pptx)
- BaseSlideExtractor : abstract interface; new adapters (e.g. AMI) must subclass this
"""

from __future__ import annotations

import re
import time
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from src.utils.logging import get_logger

logger = get_logger("ppt.extract")


# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------

@dataclass
class TableData:
    """Rows × columns of a single table extracted from a slide."""
    rows: list[list[str]]


@dataclass
class SlideData:
    """All information extracted from a single slide."""
    slide_id: str                          # e.g. "slide_01"
    slide_number: int
    title: str | None
    text: list[str]                        # non-title text blocks (reading order)
    tables: list[TableData]
    notes: str | None
    metadata: dict[str, Any]              # populated by features extractor


@dataclass
class PresentationData:
    """Top-level container for an entire presentation."""
    meeting_id: str
    source_file: str
    num_slides: int
    slides: list[SlideData] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Abstract base extractor
# ---------------------------------------------------------------------------

class BaseSlideExtractor(ABC):
    """
    Interface every slide-input adapter must implement.

    Concrete subclasses:
      - PPTXExtractor   : standard .pptx files via python-pptx
      - (future) AMISlideAdapter : AMI Meeting Corpus projection/OCR data
    """

    @abstractmethod
    def validate(self) -> None:
        """Raise ValueError / FileNotFoundError if the source is invalid."""

    @abstractmethod
    def extract(self) -> PresentationData:
        """Return a fully populated PresentationData."""


# ---------------------------------------------------------------------------
# PPTX extractor
# ---------------------------------------------------------------------------

class PPTXExtractor(BaseSlideExtractor):
    """
    Extract slide content from a standard .pptx file using python-pptx.

    Parameters
    ----------
    pptx_path : Path | str
        Path to the input .pptx file.
    meeting_id : str | None
        Optional meeting identifier.  If None, derived from the filename.
    extract_text : bool
        Whether to extract text content (default True).
    extract_tables : bool
        Whether to extract table contents (default True).
    extract_notes : bool
        Whether to attempt speaker-notes extraction (default True).
    """

    def __init__(
        self,
        pptx_path: Path | str,
        meeting_id: str | None = None,
        *,
        extract_text: bool = True,
        extract_tables: bool = True,
        extract_notes: bool = True,
    ) -> None:
        self.pptx_path = Path(pptx_path)
        self.meeting_id = meeting_id or _safe_id(self.pptx_path.stem)
        self.extract_text = extract_text
        self.extract_tables = extract_tables
        self.extract_notes = extract_notes

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def validate(self) -> None:
        if not self.pptx_path.exists():
            raise FileNotFoundError(f"PPT file not found: {self.pptx_path}")
        if self.pptx_path.suffix.lower() not in {".ppt", ".pptx"}:
            raise ValueError(
                f"Unsupported file type '{self.pptx_path.suffix}'. "
                "Only .pptx files are fully supported."
            )

    def extract(self) -> PresentationData:
        try:
            from pptx import Presentation  # type: ignore
        except ImportError as exc:
            raise ImportError("python-pptx is required: pip install python-pptx") from exc

        self.validate()

        logger.info("Opening presentation: %s", self.pptx_path)
        t0 = time.time()

        prs = Presentation(str(self.pptx_path))
        num_slides = len(prs.slides)
        logger.info("Loaded %d slides", num_slides)

        slides: list[SlideData] = []
        total_text_slides = 0
        total_tables = 0
        total_images = 0

        for idx, slide in enumerate(prs.slides, start=1):
            slide_id = f"slide_{idx:02d}"
            title = self._get_title(slide)
            text_blocks = self._get_text_blocks(slide, title) if self.extract_text else []
            tables = self._get_tables(slide) if self.extract_tables else []
            notes = self._get_notes(slide) if self.extract_notes else None
            metadata = self._get_metadata(slide, prs)

            if text_blocks or title:
                total_text_slides += 1
            total_tables += len(tables)
            total_images += metadata.get("image_count", 0)

            slides.append(
                SlideData(
                    slide_id=slide_id,
                    slide_number=idx,
                    title=title,
                    text=text_blocks,
                    tables=tables,
                    notes=notes,
                    metadata=metadata,
                )
            )

        elapsed = time.time() - t0
        logger.info(
            "Extraction complete in %.2fs | slides_with_text=%d | tables=%d | images=%d",
            elapsed,
            total_text_slides,
            total_tables,
            total_images,
        )

        return PresentationData(
            meeting_id=self.meeting_id,
            source_file=str(self.pptx_path),
            num_slides=num_slides,
            slides=slides,
        )

    # ------------------------------------------------------------------
    # Private helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _get_title(slide: Any) -> str | None:
        """Return the slide title from the title placeholder, if present."""
        try:
            from pptx.util import Pt  # noqa: F401 – just confirm library available
        except ImportError:
            pass

        # python-pptx exposes slide.shapes.title
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame:
            text = title_shape.text_frame.text.strip()
            return text or None
        return None

    @staticmethod
    def _get_text_blocks(slide: Any, title_text: str | None) -> list[str]:
        """
        Return text content from all text-bearing shapes, in reading order
        (top-to-bottom, left-to-right by bounding-box top coordinate).
        Title text is excluded to avoid duplication.
        """
        from pptx.enum.shapes import PP_PLACEHOLDER  # type: ignore

        blocks: list[tuple[int, int, str]] = []  # (top, left, text)

        for shape in slide.shapes:
            # Skip the title placeholder — already captured
            if shape.is_placeholder:
                ph_type = shape.placeholder_format.type
                if ph_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    continue

            if not shape.has_text_frame:
                continue

            raw = shape.text_frame.text.strip()
            if not raw:
                continue

            # Skip if the text is exactly the title (belt-and-suspenders)
            if title_text and raw == title_text:
                continue

            top = shape.top if shape.top is not None else 0
            left = shape.left if shape.left is not None else 0
            blocks.append((top, left, raw))

        # Sort by vertical position, then horizontal
        blocks.sort(key=lambda b: (b[0], b[1]))
        return [b[2] for b in blocks]

    @staticmethod
    def _get_tables(slide: Any) -> list[TableData]:
        """Extract all tables from a slide."""
        tables: list[TableData] = []
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            tbl = shape.table
            rows: list[list[str]] = []
            for row in tbl.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            tables.append(TableData(rows=rows))
        return tables

    @staticmethod
    def _get_notes(slide: Any) -> str | None:
        """
        Extract speaker notes.

        python-pptx exposes notes via slide.notes_slide.notes_text_frame.
        Returns None if no notes exist or the notes text is empty.
        """
        try:
            notes_slide = slide.notes_slide
            text = notes_slide.notes_text_frame.text.strip()
            return text if text else None
        except Exception:
            # Notes slide may not exist; treat as no notes
            return None

    @staticmethod
    def _get_metadata(slide: Any, prs: Any) -> dict[str, Any]:
        """Collect lightweight structural metadata for a slide."""
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore

        image_count = 0
        shape_count = 0
        table_count = 0
        text_box_count = 0
        has_chart = False

        for shape in slide.shapes:
            shape_count += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
            if shape.has_table:
                table_count += 1
            if shape.has_text_frame and not shape.is_placeholder:
                text_box_count += 1
            if shape.has_chart:
                has_chart = True

        # Slide dimensions in centimetres (EMU / 914400)
        width_cm = round(prs.slide_width / 914400 * 2.54, 2)
        height_cm = round(prs.slide_height / 914400 * 2.54, 2)

        return {
            "image_count": image_count,
            "table_count": table_count,
            "shape_count": shape_count,
            "text_box_count": text_box_count,
            "has_chart": has_chart,
            "slide_width_cm": width_cm,
            "slide_height_cm": height_cm,
        }


# ---------------------------------------------------------------------------
# Utility
# ---------------------------------------------------------------------------

def _safe_id(name: str) -> str:
    """Convert an arbitrary filename stem into a safe snake_case meeting ID."""
    safe = re.sub(r"[^\w]+", "_", name).strip("_").lower()
    return safe or "meeting"


def derive_meeting_id(pptx_path: Path | str) -> str:
    """Public helper: derive a meeting ID from a file path."""
    return _safe_id(Path(pptx_path).stem)
