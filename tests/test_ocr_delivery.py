"""Offline OCR boundary checks; the real engine case uses explicitly provisioned models."""
import hashlib
import io
import json
from pathlib import Path
from unittest.mock import patch
import pytest


def test_ocr_provision_checks_content_and_reuses_local_files(tmp_path, monkeypatch):
    from scripts import download_ocr_models as downloader
    data = b"model-test-content"
    expected = hashlib.sha1(b"blob "+str(len(data)).encode()+b"\0"+data).hexdigest()
    monkeypatch.setattr(downloader, "BLOBS", {"eng": expected})
    with patch.object(downloader.urllib.request, "urlopen", side_effect=lambda *a, **k: io.BytesIO(data)):
        downloader.provision(tmp_path)
    with patch.object(downloader.urllib.request, "urlopen", side_effect=AssertionError("Network on cached provisioning")):
        downloader.provision(tmp_path)
    assert json.loads((tmp_path/"manifest.json").read_text())["models"]["eng"]["sha256"] == hashlib.sha256(data).hexdigest()
    monkeypatch.setattr(downloader, "BLOBS", {"eng": "0"*40})
    with patch.object(downloader.urllib.request, "urlopen", side_effect=lambda *a, **k: io.BytesIO(b"bad")):
        with pytest.raises(RuntimeError, match="verification"):
            downloader.provision(tmp_path)
    assert (tmp_path/"eng.traineddata").read_bytes() == data


def test_requested_ocr_cannot_silently_return_empty_without_models(monkeypatch):
    import src.ppt.ocr as ocr
    monkeypatch.setattr(ocr, "get", lambda key, default=None: True if key=="ppt.ocr_enabled" else default)
    monkeypatch.setattr(ocr, "_ocr_available", lambda: False)
    with pytest.raises(RuntimeError, match="before enabling"):
        ocr.extract_ocr_text("slide_01", [b"image"])


def test_real_offline_hindi_english_image_slide_reaches_embedding_text(tmp_path, monkeypatch):
    from src.utils import config
    import src.ppt.ocr as ocr
    if not all((config.PROJECT_ROOT/"models/tesseract"/(lang+".traineddata")).is_file() for lang in ["eng", "hin"]):
        pytest.skip("OCR models must be explicitly provisioned")
    if not ocr.shutil.which("tesseract"):
        pytest.skip("Tesseract is not installed")
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from pptx.util import Inches
    from src.ppt.extract import PPTXExtractor
    from src.ppt.embeddings import build_slide_text
    import subprocess
    font_path=Path(subprocess.check_output(["fc-match", "-f", "%{file}", "Noto Sans Devanagari"], text=True))
    image=Image.new("RGBA", (1400, 400), (0, 0, 0, 0))
    draw=ImageDraw.Draw(image)
    english_font=subprocess.check_output(["fc-match", "-f", "%{file}", "DejaVu Sans"], text=True)
    draw.text((45, 40), "Budget review 500", font=ImageFont.truetype(english_font, 64), fill="black")
    draw.text((45, 190), "बजट रिपोर्ट", font=ImageFont.truetype(str(font_path), 68), fill="black")
    picture=tmp_path/"slide.png"; image.save(picture)
    deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6]); slide.shapes.add_picture(str(picture), Inches(0), Inches(0), width=Inches(10))
    target=tmp_path/"image-only.pptx"; deck.save(target)
    monkeypatch.setitem(config.load_config()["ppt"], "ocr_enabled", True)
    with patch("socket.socket.connect", side_effect=AssertionError("OCR attempted network access")):
        result=PPTXExtractor(target).extract()
    text=build_slide_text(result.slides[0])
    assert "Budget" in text and "500" in text and "बजट" in text
    assert result.slides[0].metadata["ocr"][0]["source"] == "ocr"
    assert result.slides[0].metadata["ocr"][0]["confidence"] is None
