"""Build the 15-minute ESE deck, editable diagrams and rehearsal material."""
import base64
import hashlib
import html
import json
import math
import re
import shutil
import textwrap
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
HERE=Path(__file__).resolve().parent
ASSETS=HERE/'assets'; ASSETS.mkdir(exist_ok=True)
EV=ROOT/'professor_guidelines/evaluation'
load=lambda p:json.loads((EV/p).read_text())
QUALITY=load('quality_scorecard/summary.json')
COMPARISON=load('classifier_comparison/summary.json')
PIPELINE=load('pipeline_improvement/summary.json')
AUDIO=load('comparative_evaluation.json')['pooled_audio_comparison']
FONT='Noto Sans'
FONT_FILE='/usr/share/fonts/noto/NotoSans-Regular.ttf'
BOLD_FILE='/usr/share/fonts/noto/NotoSans-Bold.ttf'
C={'ink':'152C40','muted':'526675','red':'B30D2C','teal':'087F8C','gold':'AF6A18',
   'paper':'FFFFFF','soft':'F0F4F6','rose':'FBEEF1','aqua':'E9F5F4','amber':'FFF3E2','line':'D9E2E7'}
W,H=13.333333,7.5
TEAM=['Anushka Chatterjee','Amitkumar Racha','Sahil Awatramani','Shriyans Sahoo']
TIMES=[30,45,50,40,35,50,45,90,45,50,70,60,45,80,65,30,10,0]
OWNERS=[TEAM[0]]*4+[TEAM[1]]*4+[TEAM[2]]*4+[TEAM[3]]*6
SLIDES=[];ELEMENTS=[]


def rgb(c):return RGBColor.from_string(C.get(c,c))
def esc(t):return html.escape(str(t))


def wrapped(text,width,size,bold=False,family=FONT):
    path='/usr/share/fonts/noto/NotoSerif-'+('Bold' if bold else 'Regular')+'.ttf' if family=='Noto Serif' else (BOLD_FILE if bold else FONT_FILE)
    font=ImageFont.truetype(path,round(size*96/72))
    max_px=width*96
    result=[]
    for original in str(text).split('\n'):
        words=original.split();line=''
        for word in words:
            candidate=(line+' '+word).strip()
            if line and font.getlength(candidate)>max_px:result.append(line);line=word
            else:line=candidate
        result.append(line)
    return result


class Deck:
    def __init__(self):
        self.p=Presentation();self.p.slide_width=Inches(W);self.p.slide_height=Inches(H)
        self.spec=[];self.s=None;self.svg=[]
    def base(self,title,kicker='THE PROJECT',subtitle='',backup=False):
        self.s=self.p.slides.add_slide(self.p.slide_layouts[6]);self.svg=[]
        self.current_title=title
        self.rect(0,0,W,H,'paper',line=None)
        self.rect(.5,.42,.065,.36,'red',line=None)
        self.text(kicker,.72,.35,10,.25,11,'red',True)
        self.text(title,.53,.82,11.9,.55,29,'ink',True,minsize=24)
        if subtitle:self.text(subtitle,.55,1.42,12.1,.45,15,'muted')
        self.s.shapes.add_picture(str(ASSETS/'image3.png'), Inches(11.15), Inches(.35),width=Inches(1.65))
        self.line(.55,7.02,12.78,7.02,'line',1,False)
        self.text('SIT Pune  |  AI & ML  |  Group 12  |  Prepared 20 October 2026',.55,7.13,10.4,.19,9,'muted')
        idx=len(self.p.slides)
        self.text(('B'+str(idx-18) if backup else f'{idx:02d} / 18'),11.8,7.12,1,.22,10,'muted',align='right')
        return self.s
    def rect(self,x,y,w,h,fill='soft',line='line',radius=False):
        s=self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
        if line:s.line.color.rgb=rgb(line);s.line.width=Pt(.8)
        else:s.line.fill.background()
        if radius:
            try:s.adjustments[0]=.06
            except Exception:pass
        self.svg.append(f'<rect x="{x*120}" y="{y*120}" width="{w*120}" height="{h*120}" rx="{8 if radius else 0}" fill="#{C.get(fill,fill)}" stroke="#{C.get(line,line) if line else C.get(fill,fill)}"/>')
        return s
    def text(self,text,x,y,w,h,size=22,color='ink',bold=False,align='left',minsize=None,family=FONT):
        minsize=size if minsize is None else minsize
        while True:
            lines=wrapped(text,w-.03,size,bold,family)
            if size*(1.08+1.18*(len(lines)-1))/72<=h+.005:break
            size-=.5
            assert size>=minsize, f'Text needs more room: {text[:65]} ({w}x{h})'
        s=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame
        tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
            p.font.name=family;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(size*1.18)
        if not hasattr(s,'name'):pass
        s.name='TEXT '+str(text)[:40]
        ELEMENTS.append({'slide':len(self.p.slides),'text':str(text),'x':x,'y':y,'w':w,'h':h,'font':size})
        for i,line in enumerate(lines):
            tx=x+w/2 if align=='center' else (x+w if align=='right' else x)
            anchor={'left':'start','right':'end','center':'middle'}[align]
            self.svg.append(f'<text x="{tx*120}" y="{(y+(i+1)*size*1.18/72)*120}" font-family="Noto Sans,sans-serif" font-size="{size*120/72}" font-weight="{700 if bold else 400}" text-anchor="{anchor}" fill="#{C.get(color,color)}">{esc(line)}</text>')
        return s
    def line(self,x1,y1,x2,y2,color='muted',width=1.5,arrow=True,dash=False):
        s=self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        s.line.color.rgb=rgb(color);s.line.width=Pt(width)
        if arrow:
            e=OxmlElement('a:tailEnd');e.set('type','triangle');s._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(e)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        self.svg.append(f'<path d="M {x1*120} {y1*120} L {x2*120} {y2*120}" fill="none" stroke="#{C.get(color,color)}" stroke-width="{width*1.6}"'+(' marker-end="url(#arrow)"' if arrow else '')+(' stroke-dasharray="6 4"' if dash else '')+'/>')
    def path(self,points,color='muted',dash=False):
        for i,(a,b) in enumerate(zip(points[:-1],points[1:])):self.line(*a,*b,color=color,arrow=i==len(points)-2,dash=dash)
    def card(self,x,y,w,h,title,body,fill='soft',size=22):
        self.rect(x,y,w,h,fill,line=None,radius=True)
        self.text(title,x+.22,y+.2,w-.44,.55,21,'ink',True)
        self.text(body,x+.22,y+.92,w-.44,h-1.06,size)
    def tag(self,text,x,y,w,fill='aqua',color='teal'):
        self.rect(x,y,w,.4,fill,line=None,radius=True);self.text(text,x+.1,y+.085,w-.2,.25,11.5,color,True)
    def note(self,key,script,bridge='',evidence='',cue=''):
        n=len(self.p.slides);main=n<=18
        owner=OWNERS[n-1] if main else 'Any team member during questions'
        seconds=TIMES[n-1] if main else 0
        notes=f'PRESENTER: {owner}\nTIME: {seconds} seconds'+(' — outside the 15-minute core; open only if asked' if not main else '')
        notes+=f'\n\nONE MESSAGE: {key}\n\nSAY:\n{script}'
        if cue:notes+='\n\nPOINT / DEMO:\n'+cue
        if bridge:notes+='\n\nTRANSITION:\n'+bridge
        if evidence:notes+='\n\nEVIDENCE / CAVEAT:\n'+evidence
        self.s.notes_slide.notes_text_frame.text=notes
        self.spec.append({'slide':n,'title':self.current_title,
            'owner':owner,'seconds':seconds,'key':key,'script':script,'bridge':bridge,'evidence':evidence,'cue':cue,'backup':not main})
    def source(self,text):self.text(text,.58,6.65,12.1,.26,10.5,'muted')
    def table(self,headers,rows,x,y,width,rowh=.43,fontsize=17,widths=None):
        widths=widths or [width/len(headers)]*len(headers)
        for ri,row in enumerate([headers]+rows):
            xx=x
            for ci,(value,ww) in enumerate(zip(row,widths)):
                self.rect(xx,y+ri*rowh,ww,rowh,'ink' if ri==0 else ('soft' if ri%2 else 'paper'),line='paper')
                self.text(str(value),xx+.1,y+ri*rowh+.075,ww-.2,rowh-.08,fontsize,'paper' if ri==0 else 'ink',bold=ri==0)
                xx+=ww
    def save(self,path):self.p.save(path)


def node(d,x,y,w,h,title,detail='',fill='soft'):
    d.rect(x,y,w,h,fill,line='line',radius=True)
    d.text(title,x+.12,y+.1,w-.24,.42,13.5,'ink',True)
    if detail:d.text(detail,x+.12,y+.53,w-.24,h-.56,11.2,'muted')


def architecture(d,title='The complete local architecture'):
    d.base(title,'SYSTEM DESIGN','Data flow, model boundaries, persistence and human review — all on the laptop.')
    for x,w,title in [(.55,2.1,'1  Capture'),(2.95,3.25,'2  Speech & slides'),(6.55,3.1,'3  Analysis'),(10.0,2.75,'4  Save & review')]:
        d.text(title,x,1.95,w,.4,17,'teal',True)
    # Connectors first: native lines remain behind all module labels.
    d.path([(1.6,3.14),(1.6,3.36)])
    d.path([(1.6,4.18),(1.6,4.4)])
    d.path([(2.65,4.85),(2.8,4.85),(2.8,2.82),(2.95,2.82)])
    d.path([(2.8,2.82),(2.8,3.96),(2.95,3.96)])
    d.path([(4.48,2.82),(4.65,2.82)])
    d.path([(4.48,3.96),(4.65,3.96)])
    d.path([(6.18,2.82),(6.27,2.82),(6.27,4.9),(6.18,4.9)],'teal')
    d.path([(5.4,4.4),(5.4,4.62)])
    d.path([(2.65,5.95),(2.95,5.95)])
    d.path([(4.48,5.95),(4.65,5.95)])
    d.path([(5.4,5.4),(5.4,5.57)])
    d.path([(6.18,5.02),(6.37,5.02),(6.37,2.85),(6.55,2.85)],'teal')
    d.path([(6.18,5.95),(6.45,5.95),(6.45,4.03),(6.55,4.03)])
    d.path([(8.08,4.42),(8.08,4.65)])
    d.path([(8.08,5.45),(8.08,5.67)])
    d.path([(9.65,2.9),(9.8,2.9),(9.8,3.0),(10.0,3.0)],'red')
    d.path([(9.65,6.06),(9.87,6.06),(9.87,3.45),(10.0,3.45)])
    d.path([(11.37,3.65),(11.37,3.9)])
    d.path([(11.37,4.77),(11.37,5.03)])
    d.path([(12.75,5.5),(12.93,5.5),(12.93,4.33),(12.75,4.33)],'teal',True)
    node(d,.55,2.45,2.1,.69,'Browser input')
    node(d,.55,3.36,2.1,.82,'FastAPI','Auth + validation')
    node(d,.55,4.4,2.1,.9,'Queue + FFmpeg','1 worker; 16 kHz mono')
    node(d,.55,5.57,2.1,.78,'Optional PPTX','Ordered slide files')
    node(d,2.95,2.45,1.53,.75,'ASR + VAD','faster-whisper','aqua')
    node(d,4.65,2.45,1.53,.75,'Word cache','Words + times','aqua')
    node(d,2.95,3.55,1.53,.85,'Diarization','Sherpa-ONNX','aqua')
    node(d,4.65,3.55,1.53,.85,'Speakers','Who spoke when','aqua')
    node(d,4.65,4.62,1.53,.78,'Attribution','Tagged words','aqua')
    node(d,2.95,5.57,1.53,.78,'Slide parsing','Text/tables/notes')
    node(d,4.65,5.57,1.53,.78,'Align or skip','MiniLM / TF-IDF')
    node(d,6.55,2.45,3.1,1.05,'Learned English review','TF-IDF → 75% forest + 25% linear\nSmooth → threshold → suggestions','rose')
    node(d,6.55,3.65,3.1,.77,'Event rules → evidence','Proposals / questions / agreements')
    node(d,6.55,4.65,3.1,.8,'Decisions → interactions','Source anchors and relationships')
    node(d,6.55,5.67,3.1,.78,'Exploratory influence','Unvalidated heuristic','amber')
    node(d,10.0,2.45,2.75,1.2,'Report builder','Suggestions + rule outputs\nSource IDs stay attached','aqua')
    node(d,10.0,3.9,2.75,.87,'Local persistence','SQLite reports + file artifacts')
    node(d,10.0,5.03,2.75,1.3,'Review + exports','Listen • labels • evidence\nJSON / MD / HTML / CSV','aqua')
    d.source('Arrows show data dependencies, not parallel compute. ASR drafts precede final speaker attribution; cached words are reused.')
    return '\n'.join(d.svg)


def build():
    d=Deck()
    # 1
    d.s=d.p.slides.add_slide(d.p.slide_layouts[6]);d.svg=[]
    d.current_title='MAI Project AY 2026–27'
    d.rect(0,0,W,H,'paper',line=None)
    d.text('MAI Project AY 2026–27',.6,.25,12.13,.52,28,'000000',True,'center',family='Noto Serif')
    d.text('ESE Presentation',.6,.93,12.13,.48,25,'000000',True,'center',family='Noto Serif')
    d.text('Project ID: Group 12',.6,1.55,12.13,.36,19,'000000',align='center',family='Noto Serif')
    d.text('Title of the Project: InfluenceIQ — Offline Decision-Discussion Detection with\nSource-Linked Meeting Review',.7,2.02,11.93,.68,19,'000000',align='center',family='Noto Serif')
    d.s.shapes.add_picture(str(ASSETS/'image1.png'),Inches(4.47),Inches(2.82),width=Inches(4.4))
    d.text('Name of the Guide: Dr. Nivedita Mishra',.6,4.9,12.13,.4,20,'000000',align='center',family='Noto Serif')
    d.text('Group Members:',.55,5.48,12.2,.34,18,'000000',family='Noto Serif')
    for i,(name,roll) in enumerate(zip(TEAM,['23070126015','24070126501','23070126112','23070126125'])):
        d.text(f'{i+1}. {name} ({roll})',.72,5.93+i*.31,11.9,.3,17.5,'000000',family='Noto Serif')
    d.text('20-10-2026',.6,7.25,2,.18,9,'muted')
    d.text('Department of Artificial Intelligence & Machine Learning',3.1,7.25,9.6,.18,9,'muted',align='center')
    d.note('MAI project: local meeting review with source evidence.',
      'Good evening. We are Group 12, presenting our MAI project, InfluenceIQ, under the guidance of Dr. Nivedita Mishra at Symbiosis Institute of Technology, Pune. Our aim is to help a reviewer find relevant decision discussion and return to the original words, speaker and recording time. We will explain the full architecture, compare the learned discussion detector with other classifiers, and show the scope of the results.',
      'First, why is a meeting recording alone difficult to review?',evidence='Cover follows the supplied reference layout with the corrected MAI Project label. Team contacts: Anushka — anushkac2311@gmail.com; Amitkumar — rachaamitkumar@gmail.com; Sahil — sahilawatramani@gmail.com; Shriyans — shriyanssahoo2005@gmail.com. Faculty guide: nivedita.mishra@sitpune.edu.in.')
    # 2
    d.base('Introduction','WHY THIS PROJECT')
    d.text('“Where did we discuss the prototype decision?”',.7,1.95,11.9,.85,30,'ink',True)
    for i,(title,body) in enumerate([('Find the passage','Relevant discussion can be scattered across the recording.'),('Know who spoke','Wrong speaker labels can credit the wrong participant.'),('Check the source','A candidate should lead back to its words and timestamp.')]):
        d.card(.65+i*4.2,3.2,3.95,2.65,title,body,size=22)
    d.source('Problem framing: passage retrieval and evidence review. We do not assume any external product lacks these features.')
    d.note('Our target is a review task people can understand.',
      'A recording contains the whole meeting, but a reviewer still has to locate the part that matters. There are three practical questions: where was the topic discussed, who spoke, and what exactly was said? Errors in any of these can change the interpretation. We therefore designed a review workflow around source passages. The system helps a person navigate and inspect evidence. It does not replace the person who decides whether a proposal was actually accepted.',
      'A small example makes that distinction clearer.')
    # 3
    d.base('Problem Statement','A PLAIN-LANGUAGE EXAMPLE','Illustrative dialogue only — these are not measured model predictions.')
    d.card(.65,2.15,5.75,2.75,'Discussion','“Could we use the smaller prototype?\nLet’s check its battery life first.”','aqua',25)
    d.card(6.75,2.15,5.9,2.75,'Final-decision candidate','“We agree to use the smaller prototype.”','rose',25)
    d.text('The learned model suggests discussion passages.\nDecision rules and human confirmation are separate.',.85,5.45,11.7,.95,26,'ink',True)
    d.note('A discussion label does not prove that a final decision was made.',
      'The statement on the left discusses a possible choice and asks for evidence. It is relevant to a decision, but the choice is not final. The statement on the right sounds like an agreement and can become a candidate for further review. Our supervised model is trained for the broader discussion-membership task. Its output is therefore a passage to inspect, not an automatic final decision. These sentences are illustrative examples; we are not presenting them as a model-accuracy test.',
      'That gives us a precise set of project objectives.')
    # 4
    d.base('Objectives of the Project','OBJECTIVES & REQUIREMENTS')
    for i,(title,body) in enumerate([('Process locally','Audio, speaker analysis and available slide content are processed on the laptop.'),('Keep evidence attached','Retain source text, segment IDs, speaker labels and recording timestamps.'),('Measure specific tasks','Compare discussion classifiers and speech components with matched references.')]):
        d.card(.65+i*4.2,2.1,3.95,3.2,title,body,size=22)
    d.text('Required behaviour: record/upload → analyse → review → export',.8,5.8,11.8,.6,25,'teal',True)
    d.note('Local processing, traceability and task-specific evaluation are the design requirements.',
      'Our first objective is local processing after model provisioning. The second is traceability: a reviewer should see the original text and jump to its time in the recording. The third is evaluation with a clearly defined task. We compare models using the same discussion references, and evaluate speech errors separately. Functional requirements include audio upload or recording, optional presentations, editable participant labels, explicit stage status and source-preserving exports.',
      'We positioned those objectives against related work rather than copying unrelated benchmark numbers.')
    # 5
    d.base('Project Plan with Timeline','PROJECT PLAN • COMPLETED PHASES','Phase sequence is documented; unverified month-by-month dates are not invented.')
    phases=[('1','Scope + literature'),('2','Prepare grouped data'),('3','Train + compare'),('4','Select combination'),('5','Integrate review'),('6','Evaluate + report')]
    for i,(n,t) in enumerate(phases):
        col=i%3;row=i//3;x=.65+4.2*col;y=2.15+1.8*row
        d.rect(x,y,3.95,1.35,'soft',line=None,radius=True);d.text(n,x+.2,y+.2,.5,.5,30,'red',True);d.text(t,x+.85,y+.25,2.85,.9,22,'ink',True)
    d.note('All six phases have artifacts; a fresh independent final evaluation is still future work.',
      'The project progressed through six documented phases: defining scope, preparing project-separated data, comparing classifiers, selecting a score combination, integrating source-based review, and recording evaluation artifacts. We have removed the rough deck’s empty month fields. There is no evidence for precise phase dates, so this shows the completed sequence honestly. Independent final evaluation is not marked complete.',
      'Here is what a user actually reviews.')
    # 6
    d.base('Literature Review','LITERATURE & CONTRIBUTION')
    rows=[['CALO [1]','Integrated meeting assistance','Bound the task to local evidence review'],['Whisper / WhisperX [2–3]','Recognition and alignment','Use pretrained speech; no ASR novelty claim'],['Speakers [4–5, 7–8]','Speakers and multilingual evaluation','Measure speaker and language limits separately'],['Meeting datasets [10–11]','Summaries and supporting evidence','Provide extractive, source-linked suggestions']]
    d.table(['Research area','Established focus','Our implemented focus'],rows,.65,2.12,12.0,rowh=.73,fontsize=17,widths=[3.0,4.0,5.0])
    d.text('Integrated local review + a matched discussion-classifier comparison.',.8,6.03,11.7,.5,23,'teal',True)
    d.note('The contribution is bounded engineering and comparative evidence, not a new speech model.',
      'CALO established a broad meeting-assistant architecture. Whisper and alignment work provide the speech context, and diarization benchmarks show why speaker errors deserve a separate measure. Meeting datasets also study summarisation and evidence. Our implemented focus is narrower: local processing with extractive review suggestions, source links and a matched discussion-classifier comparison. We do not compare our F1 score with a summarisation score or claim that every external system is cloud-only or untraceable.',
      'The next slide defines the project gap before we examine the complete solution.',evidence='References follow the paper. This is a scope comparison, not a head-to-head benchmark of complete assistants.')
    # 7
    d.base('Gap in the Research / Technology / Methodology','PROJECT FOCUS','Established components exist; the project connects them into a reviewable local workflow.')
    for i,(title,body) in enumerate([('Integration need','Connect words, speakers and optional slide content while retaining source IDs and timestamps.'),('Evaluation need','Measure discussion membership separately from speech errors and final-decision correctness.'),('Evidence still needed','Validate complete recordings, slide links and human influence rankings on new annotations.')]):
        d.card(.65+i*4.2,2.1,3.95,3.65,title,body,'soft',22)
    d.text('Implemented focus: local, source-linked discussion review.',.8,6.05,11.9,.45,24,'teal',True)
    d.note('The gap is the bounded integration and evidence needed for this project.',
      'Related work already provides recognition, diarization and meeting-assistant ideas. We therefore do not claim that every other assistant lacks local processing or source links. Our engineering task is to connect these components and preserve traceability. Our evaluation task is to separate discussion retrieval from speech errors and final decisions. Independent labels for the complete recording workflow, slide links and influence remain missing. This is our bounded project contribution and the next research need.',
      'The full architecture shows how these parts connect.')
    # 8
    arch_svg=architecture(d,'Description of the Proposed Solution')
    d.note('The full architecture separates data processing, learned suggestions and review.',
      'Read this diagram from left to right. First the browser supplies audio and optional slides; the local API validates inputs and queues one analysis job. Second, speech recognition supplies words and times, while speaker clustering supplies speaker intervals. Cached words are combined with the final speaker intervals. Optional slides are extracted and matched to passages. Third, the learned English discussion branch produces review suggestions. The separate event, evidence, decision and interaction branch supports other candidate outputs and exploratory influence scores. Finally, reports and local persistence feed the review interface. Human edits return to the saved application report. These are data dependencies, not a claim that the modules run simultaneously.',
      'These processing and review stages define the requirements.',evidence='Verified against backend/api.py, src/pipeline/orchestrator.py, src/audio/asr.py, src/meeting/discussions.py and report/export modules. The learned discussion stage does not create final-decision anchors or feed the influence score.',cue='Use three sweeps: capture/preparation; speech and slides; analysis and saved review. Do not read all boxes. A separate editable architecture deck and vector diagram are supplied for close inspection.')
    # 9
    d.base('Requirement Analysis','FUNCTIONAL & OPERATIONAL REQUIREMENTS')
    d.card(.65,2.1,5.9,3.9,'Functional requirements','Record or upload audio.\nOptionally add presentation slides.\nReview source text and timestamps.\nCorrect speaker display names.\nExport source-linked reports.','aqua',23)
    d.card(6.85,2.1,5.8,3.9,'Operational requirements','Preprovision local model assets.\nShow explicit processing stages.\nQueue one analysis job at a time.\nPersist job and report state.\nKeep optional-stage failures visible.','soft',23)
    d.source('Speaker renaming changes display metadata; correcting acoustic clusters requires speaker analysis to be rerun.')
    d.note('Requirements cover the user journey and processing reliability.',
      'The user must be able to supply audio, optionally attach slides, inspect source passages, correct speaker names and export a report. Operationally, the local models must already be provisioned, processing stages must be visible and the worker must queue analysis. Job state and report artifacts support recovery where available. An optional model failure receives an explicit status. Renaming a speaker changes the displayed identity; repairing a wrong acoustic cluster requires diarization and attribution to run again, reusing unchanged word observations.',
      'The technology stack maps these requirements to concrete components.')
    # 10
    d.base('Technology Stack','IMPLEMENTED COMPONENTS')
    d.table(['Layer','Technology','Purpose'],[
        ['Interface','React + browser audio','Record, upload and source review'],
        ['Local service','FastAPI + Python','Validation, queue and orchestration'],
        ['Audio and speakers','FFmpeg / faster-whisper / Sherpa-ONNX','Normalize, recognise and diarize'],
        ['Slides','PPTX parsing / MiniLM / TF-IDF','Extract text and propose links'],
        ['Learned discussion','scikit-learn: TF-IDF + forest / logistic','English discussion suggestions'],
        ['Storage and reports','SQLite + local artifacts','Saved review and report exports']],.65,2.05,12.0,.64,17,[2.55,5.0,4.45])
    d.source('Speech and embedding models are pretrained. Small supervised discussion classifiers are trained on the training partition.')
    d.note('The project trains lightweight classifiers while reusing pretrained speech and embeddings.',
      'The browser interface uses React, while FastAPI and Python coordinate the local workflow. FFmpeg normalizes the audio; faster-whisper recognizes words; Sherpa-ONNX supplies speaker segmentation and clustering. Presentation parsing extracts native text, tables and notes, and text similarity can propose slide links. The supervised English detector uses train-fitted lexical features and conventional classifiers. SQLite and local artifacts retain reports and job state. We do not fine-tune a speech or embedding foundation model. We do train small classifiers for the discussion task.',
      'Here is the design of that trained NLP component.')
    # 11
    d.base('Design: Discussion Detector','THE NLP METHOD','Two conventional classifiers; no generative meeting summary is produced.')
    stages=[('Reference acts','Human text +\ndiscussion labels'),('Lexical features','Words / characters\n+ nearby context'),('Two scores','Random forest\n+ logistic regression'),('Review output','Smooth scores\n+ apply threshold')]
    for i,(title,body) in enumerate(stages):
        x=.65+i*3.18;d.card(x,2.15,2.95,2.45,title,body,'soft',20)
        if i<3:d.line(x+2.98,3.35,x+3.13,3.35,'teal')
    d.rect(.8,5.05,11.7,.8,'rose',line=None,radius=True)
    d.text('Combined score = 75% forest + 25% logistic',1.05,5.22,11.1,.5,26,'red',True)
    d.text('Context: ±40 acts. Smoothing: ±5 acts. Threshold: 0.357740. English-only learned review.',.85,6.12,11.7,.4,16,'muted')
    d.note('The learned component is small supervised NLP training over fixed lexical features.',
      'The research unit is a dialogue act: a short manually delimited utterance. An act is positive when at least half its lexical word identifiers fall inside an annotated decision-discussion span. We convert words and character patterns into numerical features, and add context from up to forty acts on each side without crossing meeting boundaries. A random forest and logistic classifier each produce a score. The selected combination is seventy-five percent forest and twenty-five percent logistic, followed by smoothing and a fixed threshold. Labels and outcome categories are not input features. Future context makes this an offline detector. The speech models are pretrained, while these small classifiers are trained on our training partition.',
      'The implementation also needs a careful separation of data.')
    # 12
    d.base('Development / Implementation','EVALUATION DESIGN','47 AMI meetings with manual discussion spans; series remain in one partition.')
    d.table(['Partition','Meetings','Acts','Positive acts'],[['Training','29','20,339','5,862'],['Development','6','3,739','1,117'],['Historical','12','11,136','2,481']],.8,2.18,11.7,rowh=.72,fontsize=23,widths=[3.3,2.1,3.0,3.3])
    d.text('Compare 4 classifier families × 2 feature settings.\nChoose settings on development data; retain all results.',.85,5.35,11.6,.95,25,'ink',True)
    d.source('Historical data has been examined across rounds: exploratory evidence, with only 3 independent historical project series.')
    d.note('The cohort is grouped, but repeated historical inspection limits the final claim.',
      'We have forty-seven AMI meetings with manual discussion spans. The custom split uses twenty-nine training meetings, six development meetings and twelve historical meetings. All meetings in a project series stay together. We compare four classifier families with local or wider-context features and select each configuration on development F1. However, the historical cohort has been examined across successive rounds and has only three independent project series. We therefore label the results exploratory. Human transcripts and act boundaries are used here, so this is not the accuracy of the complete audio application.',
      'Here are the actual discussion results under that protocol.')
    # 13
    d.base('Testing & Debugging','TESTING & QUALITY')
    d.card(.65,2.1,5.9,3.7,'597 recorded software tests','Checks include batching, source IDs, meeting boundaries, optional-stage errors and report/export integration.','aqua',24)
    d.card(6.85,2.1,5.8,3.7,'Keep failed experiments visible','Semantic and continuity searches did not replace the baseline. Calibration and refit gains failed transfer checks.','soft',24)
    d.source('597-test regression recorded 19 October 2026; it is not a count of correctly analysed meetings. Full negative results remain in the research report.')
    d.note('Functional checks establish execution; independent references establish accuracy.',
      'The recorded regression passed five hundred and ninety-seven software tests. Those checks cover things such as preserved source identifiers, batch scoring, meeting boundaries, exports and optional-stage failures. They do not mean five hundred and ninety-seven correctly understood meetings. We also retained changes that did not work: semantic and continuity variants failed the development gate, while calibration and refit results did not justify replacing the installed model. A two-excerpt speaker-calibration pilot reduced diarization error, but that is another small, separately scoped task.',
      'Now let us examine the measured discussion results.',evidence='Regression log: professor_guidelines/evaluation/quality_scorecard/regression_tests.log. No new model training or application regression run is implied by building this deck.')
    # 14
    d.base('Project Outcome: Discussion Detection','DISCUSSION RESULTS','Historical cohort: human reference transcripts and act boundaries — not recording-level app accuracy.')
    values=[('Logistic',50.60),('Naive Bayes',47.12),('Linear SVM',50.48),('Forest',52.08),('Combination',53.54)]
    for i,(name,value) in enumerate(values):
        yy=2.2+i*.68;d.text(name,.75,yy,2.1,.4,18,'ink',bold=i==4)
        d.rect(3.0,yy+.05,4.2*value/65,.33,'red' if i==4 else 'teal',line=None)
        d.text(f'{value:.2f}',7.35,yy-.01,.8,.4,18,'ink',True)
    d.text('Historical F1 (%) — common context-40 features',.8,5.95,7.4,.4,16,'muted')
    for i,(value,label,detail) in enumerate([('39.94%','Precision','About 40 of 100 flagged acts\nmatch reference positives.'),('81.18%','Recall','Finds about 81 of 100\nreference-positive acts.')]):
        y=2.25+i*1.9;d.rect(8.85,y,3.65,1.65,'soft',line=None,radius=True);d.text(value,9.05,y+.12,3.2,.55,31,'red',True);d.text(label,9.05,y+.72,3.2,.35,17,'ink',True);d.text(detail,9.05,y+1.1,3.2,.5,12.5,'muted')
    d.source('Mixture F1 53.54%; matched forest 52.08%. Difference +1.46 points; descriptive series-bootstrap interval −7.16 to +5.99.')
    d.note('There is a measured F1 gain, but it is not proof of reliable or universal superiority.',
      'The chart compares the context versions on identical reference acts. The combined detector has a fifty-three point five-four percent F1 score, versus fifty-two point zero-eight for the matched forest. The gain is one point four-six percentage points. Precision is only thirty-nine point nine-four percent: roughly forty of every hundred flagged acts match the reference positives. Recall is eighty-one point one-eight percent, so we find many reference discussion acts but also flag many negatives. The descriptive uncertainty interval crosses zero, and some other classifiers have higher average precision. This is an exploratory task-level result, not a statistically established win or complete-project accuracy.',
      'Recognition quality also depends on the language and cohort.',evidence='Exact P/R/F1/AP and confusion counts are retained in the research report. Do not call 53.54% overall accuracy.')
    # 15
    d.base('Project Outcome: Speech and Speakers','SPEECH RESULTS','Word error rate: edits per 100 reference words, pooled by reference count. Lower is better.')
    rows=[('Hindi • 8 samples',50.96,28.85),('Hinglish • 8 samples',66.13,57.26),('English historical • 2 clips',61.61,63.91)]
    for i,(label,old,new) in enumerate(rows):
        y=2.15+i*1.16;d.text(label,.8,y,5.3,.45,23,'ink',True)
        d.text(f'{old:.2f}%',6.35,y,2.1,.5,28,'muted');d.line(8.5,y+.27,9.2,y+.27)
        d.text(f'{new:.2f}%',9.55,y,2.45,.5,28,'teal' if new<old else 'gold',True)
    d.text('small model',6.28,5.82,2.4,.4,17,'muted');d.text('large-v3-turbo',9.45,5.82,3.0,.4,17,'muted')
    d.text('Speaker pilot: DER 64.08% → 40.91%; speaker-count MAE 8.5 → 1.5.',.8,6.18,11.8,.35,17,'ink',True)
    d.source('Small stays the shared default; Turbo is selectable. Tiny fixed pilots, strict scoring and reference-coverage limits apply.')
    d.note('Model choice remains explicit because improvement is corpus-dependent.',
      'Shriyans here. Word error rate counts substitutions, deletions and insertions against reference words; lower is better. On eight fixed Hindi samples, Turbo lowers word error from fifty point nine-six to twenty-eight point eight-five percent. Hinglish improves from sixty-six point one-three to fifty-seven point two-six percent. But historical English becomes worse, from sixty-one point six-one to sixty-three point nine-one percent. The shared default remains small, while Turbo is available explicitly. These are tiny speech pilots, not broad multilingual meeting validation. The same unchanged samples and strict normalization were used. In the custom two-excerpt speaker pilot, diarization error fell from 64.08 to 40.91 percent after development-selected calibration. All four development and historical audio meetings are from the referenced official training partition. Segment-reference and coverage limits apply; this is not an official full-corpus benchmark.',
      'The conclusion brings the working system and remaining validation needs together.')
    # 16
    d.base('Conclusion','TAKEAWAY')
    d.text('Local processing\nSource-linked review\nMeasured, bounded claims',.9,2.05,11.5,2.6,39,'ink',True)
    d.text('A working prototype, a reproducible classifier comparison,\nand a clear plan for independent recording-level validation.',.95,5.3,11.1,1.0,25,'teal')
    d.source('Still unmeasured: final-decision, slide-link and human influence accuracy. Fresh recording-level evaluation is needed.')
    d.note('The delivered result is useful review infrastructure and honest comparative evidence.',
      'To conclude, InfluenceIQ integrates local processing with source-linked discussion review. We have a working application, a matched classifier study and transparent component results. The current discussion score is fifty-three point five-four percent F1 on oracle inputs, and the next priority is independent recording-level evaluation. Final-decision accuracy, slide-link accuracy and human influence rankings remain unmeasured. The next priority is to freeze the pipeline and annotate fresh recordings independently. Thank you.')
    # 17
    d.base('References','SELECTED REFERENCES • PAPER NUMBERING')
    paper=(ROOT/'professor_guidelines/RESEARCH_PAPER.md').read_text()
    all_refs=re.findall(r'^\[(\d+)\] (.+)$',paper.split('## References')[1],re.M)
    entries=[
        ('1','Tur et al. (2010)','The CALO meeting assistant system.'),
        ('2','Radford et al. (2023)','Robust speech recognition via large-scale weak supervision.'),
        ('3','Bain et al. (2023)','WhisperX: Time-accurate speech transcription of long-form audio.'),
        ('4','Bredin (2023)','pyannote.audio 2.1 speaker diarization pipeline: Principle, benchmark, and recipe.'),
        ('5','Plaquet & Bredin (2023)','Powerset multi-class cross entropy loss for neural speaker diarization.'),
        ('7','Baghel et al. (2023)','The DISPLACE challenge 2023.'),
        ('8','Kalluri et al. (2024)','The second DISPLACE challenge: DIarization of SPeaker and LAnguage in Conversational Environments.'),
        ('10','Hu et al. (2023)','MeetingBank: A benchmark dataset for meeting summarization.'),
        ('11','Kim et al. (2023)','ExplainMeetSum: A dataset for explainable meeting summarization aligned with human intent.'),
        ('12','Carletta et al. (2006)','The AMI meeting corpus: A pre-announcement.')]
    for i,(number,author,title) in enumerate(entries):
        x=.7+(i//5)*6.2;y=1.95+(i%5)*.91
        d.text(f'[{number}] {author}\n{title}',x,y,5.75,.88,14.5,'ink',minsize=13.5)
    d.source('Selected entries retain IEEE paper numbering. All 15 full references are in the paper and this slide’s notes.')
    d.note('Citations identify prior work; they are not complete-system benchmark comparisons.',
      'These are the main sources supporting our meeting-assistant context, speech and speaker components, and dataset discussion. The complete fifteen-entry bibliography is included in our IEEE paper and this slide’s notes. Thank you; we welcome your questions.',
      evidence='\n'.join('['+n+'] '+r for n,r in all_refs))
    # 18
    d.base('Photograph of the Presentation','MAI PROJECT • GROUP 12')
    d.rect(1.15,2.0,11.0,4.25,'soft',line='line',radius=True)
    d.text('Add the actual presentation photograph\nafter the session.',1.65,3.3,10.0,1.0,27,'muted',align='center')
    d.text('Symbiosis Institute of Technology, Pune • AY 2026–27',1.5,5.4,10.3,.45,19,'muted',align='center')
    d.note('Post-event photographic record; no photograph has been fabricated.',
      'This slide retains the heading required by the original presentation. Add an actual photograph after the presentation if the department requires a photographic record. It is not part of the spoken fifteen-minute script.',cue='For questions, return to slide 16 (Conclusion), slide 8 (architecture) or slide 14 (discussion results).')
    assert sum(TIMES)==840
    d.save(HERE/'InfluenceIQ_ESE_Final.pptx')
    architecture_deck=Deck();architecture(architecture_deck)
    for shape in architecture_deck.s.shapes:
        if shape.has_text_frame and shape.text=='01 / 18':shape.text_frame.paragraphs[0].text='ARCHITECTURE'
    architecture_deck.save(HERE/'Architecture_Detailed.pptx')
    svg='<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 1600 900"><defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="3" orient="auto"><path d="M0,0 L0,6 L7,3 z" fill="#526675"/></marker></defs>'+arch_svg+'</svg>'
    (HERE/'Architecture_Detailed.svg').write_text(svg)
    (HERE/'slide_manifest.json').write_text(json.dumps({'main_slides':18,'backup_slides':0,'spoken_seconds':840,'transition_buffer_seconds':60,'slides':d.spec},indent=2,ensure_ascii=False))
    (HERE/'layout_elements.json').write_text(json.dumps(ELEMENTS,indent=2,ensure_ascii=False))
    guide(d.spec)
    print('Built',len(d.p.slides),'slides; 14-minute script + 1-minute buffer. Original deck preserved.')


def guide(spec):
    lines=['# InfluenceIQ — 15-minute presenter guide','',
      'Use `InfluenceIQ_ESE_Final.pptx` for the presentation. The final deck contains exactly 18 slides, using the original headings and order. Slide 18 is the post-event photograph slot. The spoken plan is 14 minutes, with one minute reserved for transitions and small delays. Questions are outside the 15-minute plan. PowerPoint/Impress presenter notes contain the same prompts.',
      '', '## Suggested handoffs','',
      '| Presenter | Slides | Planned speaking time | Focus |','| --- | --- | --- | --- |']
    for name in TEAM:
        chosen=[s for s in spec[:18] if s['owner']==name];secs=sum(s['seconds'] for s in chosen)
        lines.append(f"| {name} | {chosen[0]['slide']}–{chosen[-1]['slide']} | {secs//60}:{secs%60:02d} | {chosen[0]['key']} |")
    lines+=['','These are rehearsal assignments, not statements of who implemented each feature. Swap blocks if needed.','',
      '## Nontechnical audience rules','',
      '- Explain ASR as “what was said” and diarization as “who spoke when.”',
      '- Use the illustrative discussion/decision example before the model formula.',
      '- Do not read every architecture box; follow the three sweeps in slide 8’s notes.',
      '- Define precision as “how many suggestions were right” and recall as “how much reference discussion was found.”',
      '- Say “historical discussion F1 using human text,” never “overall project accuracy.”','',
      '## Rehearsal script','']
    elapsed=0
    for s in spec[:18]:
        start=elapsed;elapsed+=s['seconds'];mm=lambda sec:f'{sec//60:02d}:{sec%60:02d}'
        lines += [f"### {s['slide']}. {s['title']}",'',f"**{s['owner']} • {mm(start)}–{mm(elapsed)} • {s['seconds']} seconds**",'',s['script']]
        if s['cue']:lines += ['', '**Point / demonstration:** '+s['cue']]
        if s['bridge']:lines += ['', '**Transition:** '+s['bridge']]
        if s['evidence']:lines += ['', '**Evidence / limit:** '+s['evidence']]
        lines += ['']
    lines += ['## Likely viva questions','',
      '**Is this training-free?** Speech and embedding models are pretrained/frozen. Small supervised discussion classifiers are trained on the training acts.',
      '', '**Why not call the result superior?** Historical F1 improves by 1.46 points over the matched forest, but the three-series descriptive interval crosses zero and AP is worse than some baselines.',
      '', '**Why is accuracy below an always-negative baseline?** The data is imbalanced. Predicting negative everywhere gives 77.72% accuracy but finds no discussion. Report positive F1, recall and balanced accuracy.',
      '', '**Can it understand Hindi decisions?** Hindi/Hinglish speech recognition has small pilots. The learned discussion model is English-only; final-decision accuracy is unmeasured.',
      '', '**Why use both forest and logistic regression?** This particular mixture was selected on development F1. We do not claim the model combination is inherently or universally better.',
      '', '**What does ±40 mean?** Up to forty dialogue acts before and after, clipped to a meeting. Future context makes the detector offline.',
      '', '**What happens when a speaker name is wrong?** Rename the displayed identity. If the acoustic grouping is wrong, rerun speaker separation; unchanged word observations can be reused.',
      '', '**Does influence measure real decision power?** No independent human ranking exists. The audited case equals scaled decision counting; no causal or validated influence claim is supported.',
      '', '**What is the novelty?** A bounded integration and comparative study: local review with source traceability, cached recognition/attribution separation, and explicit model/metric boundaries. There is no new foundation model or proven novelty over every prior system.',
      '', '**Are the interface examples real model results?** The screenshots show the real application with isolated synthetic validation data. They demonstrate UI functionality only.',
      '']
    lines += ['','## Presentation-day checklist','',
      '- Open the final PPTX and PDF fallback locally; verify fonts and the projector’s 16:9 display.',
      '- Practise at least once with all four members; stop the spoken part by 14:00 and keep the last minute as buffer.',
      '- If demonstrating, prepare a completed English recording and test the sound output. Avoid new inference or model downloads on stage.',
      '- Use Presenter View for notes; on a single screen keep this guide on a separate device.',
      '- Keep the architecture SVG/PDF ready for detailed examiner questions.',
      '- The rough deck’s “photograph of presentation” is a post-event record. Add a real photograph after the event if required; no photograph or implementation dates have been fabricated.',
      '- Team contacts are retained in the cover notes and the contact list below; speaking assignments do not imply implementation ownership.','']
    lines += ['## Team contacts', '', 'Anushka Chatterjee — anushkac2311@gmail.com', '', 'Amitkumar Racha — rachaamitkumar@gmail.com', '', 'Sahil Awatramani — sahilawatramani@gmail.com', '', 'Shriyans Sahoo — shriyanssahoo2005@gmail.com', '', 'Faculty guide: Dr. Nivedita Mishra — nivedita.mishra@sitpune.edu.in', '']
    markdown='\n'.join(lines)
    (HERE/'PRESENTER_GUIDE.md').write_text(markdown)
    from markdown_it import MarkdownIt
    content=MarkdownIt('commonmark').enable('table').render(markdown)
    (HERE/'PRESENTER_GUIDE.html').write_text('''<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>InfluenceIQ — Presenter Guide</title><style>
@page { size:A4; margin:18mm; }
body {font-family:"Noto Sans",sans-serif;font-size:11pt;line-height:1.5;color:#152C40;}
h1 {font-size:22pt;color:#B30D2C;} h2 {font-size:16pt;color:#087F8C;}
h3 {font-size:13pt;} h1,h2,h3 {page-break-after:avoid;}
table {border-collapse:collapse;width:100%;font-size:9pt;}
th,td {border:1px solid #D9E2E7;padding:6px;text-align:left;vertical-align:top;}
th {background:#F0F4F6;} a {color:#087F8C;} li {margin-bottom:5px;}
</style></head><body>'''+content+'</body></html>')


if __name__=='__main__':
    import re
    build()
