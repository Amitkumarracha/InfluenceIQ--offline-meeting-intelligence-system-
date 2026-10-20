"""Check the delivered deck, source metrics, speaker plan and rendered exports."""
import hashlib
import json
import re
import subprocess
import unicodedata
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
EV = ROOT / 'professor_guidelines/evaluation'
load = lambda p: json.loads(p.read_text())
deck = Presentation(HERE / 'InfluenceIQ_ESE_Final.pptx')
arch = Presentation(HERE / 'Architecture_Detailed.pptx')
manifest = load(HERE / 'slide_manifest.json')
quality = load(EV / 'quality_scorecard/summary.json')
comparison = load(EV / 'classifier_comparison/summary.json')
pipeline = load(EV / 'pipeline_improvement/summary.json')
audio = load(EV / 'comparative_evaluation.json')['pooled_audio_comparison']
texts = ['\n'.join(s.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
notes = [s.notes_slide.notes_text_frame.text for s in deck.slides]

assert len(deck.slides) == 18 and len(arch.slides) == 1
assert abs(deck.slide_width / deck.slide_height - 16 / 9) < .00001
assert manifest['main_slides'] == 18 and manifest['backup_slides'] == 0
assert sum(s['seconds'] for s in manifest['slides']) == 840
assert manifest['spoken_seconds'] + manifest['transition_buffer_seconds'] == 900
assert all(s['backup'] == (s['slide'] > 18) for s in manifest['slides'])
assert all(all(t in n for t in ['PRESENTER:', 'TIME:', 'ONE MESSAGE:', 'SAY:']) for n in notes)
assert 'MAI Project AY 2026–27' in texts[0]
assert not re.search(r'B\.?\s*TECH', '\n'.join(texts), re.I)
assert 'actual presentation photograph' in texts[17]
assert 'not part of the spoken' in notes[17]
assert not any('backup' in n.lower() or 'appendix' in n.lower() for n in notes)
assert not re.search(r'\b(TODO|TBD|add month|insert photo|lorem ipsum)\b', '\n'.join(texts), re.I)

members = ['Anushka Chatterjee', 'Amitkumar Racha', 'Sahil Awatramani', 'Shriyans Sahoo']
for name, roll in zip(members, ['23070126015', '24070126501', '23070126112', '23070126125']):
    assert name in texts[0] and roll in texts[0]
assert {n: sum(s['seconds'] for s in manifest['slides'] if s['owner'] == n) for n in members} == dict(zip(members, [165, 220, 225, 230]))
for email in ['anushkac2311@gmail.com', 'rachaamitkumar@gmail.com', 'sahilawatramani@gmail.com', 'shriyanssahoo2005@gmail.com', 'nivedita.mishra@sitpune.edu.in']:
    assert email in notes[0]
assert 'Name of the Guide: Dr. Nivedita Mishra' in texts[0]

def walk_text(shapes):
    for shape in shapes:
        if shape.has_text_frame and shape.text.strip():
            yield shape.text.strip()
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_text(shape.shapes)

rough = Presentation(HERE / 'InfluenceIQ_ESE_Presentation.pptx.pptx')
headings = [next(walk_text(slide.shapes)) for slide in list(rough.slides)[1:]]
assert [s['title'] for s in manifest['slides'][1:]] == headings
for title, text in zip(headings, texts[1:]):
    assert title in text

# Use the same Noto font files as the builder, but inspect the saved PPTX geometry.
font_cache = {}
native_boxes = 0
for presentation in [deck, arch]:
    for number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            assert min(shape.left, shape.top, shape.width, shape.height) >= 0, (number, shape.name)
            assert shape.left + shape.width <= presentation.slide_width + 100, (number, shape.name)
            assert shape.top + shape.height <= presentation.slide_height + 100, (number, shape.name)
            if not shape.has_text_frame or not shape.text.strip():
                continue
            native_boxes += 1
            heights = []
            for para in shape.text_frame.paragraphs:
                size = para.font.size.pt
                assert size >= 9, (number, shape.text)
                font_key = (round(size * 96 / 72), bool(para.font.bold), para.font.name)
                if font_key not in font_cache:
                    family = 'NotoSerif' if para.font.name == 'Noto Serif' else 'NotoSans'
                    font_cache[font_key] = ImageFont.truetype('/usr/share/fonts/noto/' + family + '-' + ('Bold' if font_key[1] else 'Regular') + '.ttf', font_key[0])
                available = (shape.width - shape.text_frame.margin_left - shape.text_frame.margin_right) / 914400 * 96
                assert font_cache[font_key].getlength(para.text) <= available + 3, (number, para.text)
                heights.append(size * 1.18)
            needed = sum(heights[:-1]) + para.font.size.pt * 1.08
            assert needed <= shape.height / 12700 + .5, (number, shape.text)

architecture = deck.slides[7]
assert sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in architecture.shapes) == 1
assert sum(s.shape_type == MSO_SHAPE_TYPE.LINE for s in architecture.shapes) >= 30
assert sum(s.has_text_frame and bool(s.text) for s in architecture.shapes) >= 35
for word in ['Word cache', 'Attribution', 'Diarization', 'Slide parsing', 'Learned English review', 'Local persistence', 'Exploratory influence']:
    assert word in texts[7]
assert 'not parallel compute' in texts[7] and 'does not create final-decision' in notes[7]
assert not any(s.click_action.target_slide for slide in deck.slides for s in slide.shapes)

installed = quality['discussion']['installed']
for key in ['precision', 'recall', 'f1']:
    assert f'{installed[key] * 100:.2f}' in texts[13]
for row in comparison['models']:
    if row['features'] == 'context_40':
        assert f'{row["test"]["f1"] * 100:.2f}' in texts[13]
assert pipeline['selection']['linear_weight'] == .25
assert pipeline['selection']['smoothing_radius'] == 5
assert f'{pipeline["selection"]["threshold"]:.6f}' in texts[10]
for row in quality['asr']:
    if row['model'] in ['small', 'large-v3-turbo'] and row['cohort'] != 'English development (2 clips)':
        assert f'{row["wer"]["rate"] * 100:.2f}' in texts[14]
for row in audio:
    assert f'{row["pooled_der"] * 100:.2f}' in texts[14]
    assert str(row['speaker_count_mae']) in texts[14]
assert '597 passed' in (EV / 'quality_scorecard/regression_tests.log').read_text()
assert '597' in texts[12] and '19 October 2026' in texts[12]
assert 'human reference' in texts[13].lower() and 'exploratory' in notes[13].lower()
assert 'official training partition' in notes[14]
assert all(f'[{i}]' in notes[16] for i in range(1, 16))

# Read every rendered page to catch dropped or clipped native text.
def tokens(value):
    return Counter(re.findall(r'[^\W_]+', unicodedata.normalize('NFKC', value).casefold()))

pdf_pages = {}
for filename, expected in [('InfluenceIQ_ESE_Final.pdf', 18), ('Architecture_Detailed.pdf', 1), ('PRESENTER_GUIDE.pdf', None)]:
    path = HERE / filename
    info = subprocess.check_output(['pdfinfo', str(path)], text=True)
    count = int(re.search(r'^Pages:\s+(\d+)', info, re.M).group(1))
    assert count == expected if expected is not None else count > 1
    pdf_pages[filename] = count
    xml = ET.fromstring(subprocess.check_output(['pdftotext', '-bbox-layout', str(path), '-']))
    pages = xml.findall('.//{*}page')
    for number, page in enumerate(pages):
        words = page.findall('.//{*}word')
        for word in words:
            assert float(word.get('xMin')) >= -.5 and float(word.get('yMin')) >= -.5
            assert float(word.get('xMax')) <= float(page.get('width')) + .5
            assert float(word.get('yMax')) <= float(page.get('height')) + .5
        if filename == 'InfluenceIQ_ESE_Final.pdf':
            missing = tokens(texts[number]) - tokens(' '.join(w.text or '' for w in words))
            assert not missing, (number + 1, missing)

sources = ['presentation/InfluenceIQ_ESE_Presentation.pptx.pptx', 'professor_guidelines/RESEARCH_PAPER.md'] + [str(p.relative_to(ROOT)) for p in [EV / 'quality_scorecard/summary.json', EV / 'classifier_comparison/summary.json', EV / 'pipeline_improvement/summary.json', EV / 'comparative_evaluation.json', EV / 'quality_scorecard/regression_tests.log']] + ['backend/api.py', 'src/pipeline/orchestrator.py', 'src/audio/asr.py', 'src/meeting/discussions.py']
hashes = {p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest() for p in sources}
provenance = HERE / 'source_manifest.json'
if provenance.exists():
    assert load(provenance)['sha256'] == hashes, 'Source material changed; re-review the affected slides.'
else:
    provenance.write_text(json.dumps({'prepared': '2026-10-20', 'scope': 'Presentation sources; original rough deck retained unchanged.', 'sha256': hashes}, indent=2) + '\n')

result = {'status': 'passed', 'slides': 18, 'main': 18, 'backup': 0, 'planned_seconds_including_buffer': 900, 'native_text_boxes_checked': native_boxes, 'pdf_pages': pdf_pages, 'checks': ['exactly 18 slides', 'original heading order', 'MAI cover label', 'slide and text bounds', 'rendered native text retained', 'all slide notes and handoffs', 'no removed-slide navigation links', 'classifier/ASR/diarization metric consistency', 'source hashes', 'editable architecture geometry'], 'limits': ['Slide 18 needs the actual post-event photograph.', 'Physical projector and live audio devices need a presentation-room check.', 'Recorded 597-test project run was not rerun for this document-only task.']}
(HERE / 'validation.json').write_text(json.dumps(result, indent=2) + '\n')
print(json.dumps(result, indent=2))
