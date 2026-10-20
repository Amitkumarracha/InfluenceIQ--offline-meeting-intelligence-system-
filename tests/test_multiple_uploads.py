"""Ordered multi-file uploads and complete incremental transcript delivery."""
import io
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
import soundfile as sf
from pptx import Presentation

from tests.test_api import api_client, register
from src.audio.parts import decode_parts
from src.audio.asr import run_asr
from src.ppt.parts import extract_presentations


def audio(value, rate=16000):
    stream = io.BytesIO()
    sf.write(stream, np.full(rate, value), rate, format='WAV', subtype='PCM_16')
    return stream.getvalue()


def deck(title):
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[0]).shapes.title.text = title
    stream = io.BytesIO(); presentation.save(stream)
    return stream.getvalue()


def test_ordered_audio_and_presentations_survive_upload_and_extraction(api_client, tmp_path):
    api, client = api_client
    auth = register(client, 'multi@example.com')
    files = [('audio', ('first.wav', audio(.1, 8000))),
             ('audio', ('second.wav', audio(-.2, 22050))),
             ('ppt', ('first.pptx', deck('Budget'))),
             ('ppt', ('second.pptx', deck('Delivery')))]
    response = client.post('/api/analyze', headers=auth, files=files)
    assert response.status_code == 202, response.text
    state = api.job_state(response.json()['meeting_id'])
    assert len(state['audio']) == len(state['ppt']) == 2
    assert abs(state['audio_duration_seconds'] - 2) < .01
    decoded = decode_parts(state['audio'], tmp_path / 'combined.wav')
    samples, rate = sf.read(decoded)
    assert rate == 16000 and len(samples) == 32000
    assert np.mean(samples[:16000]) > .09 and np.mean(samples[16000:]) < -.19
    slides = extract_presentations(state['ppt'], 'combined')
    assert [s.title for s in slides.slides] == ['Budget', 'Delivery']
    assert [s.slide_id for s in slides.slides] == ['slide_01', 'slide_02']
    assert [s.metadata['source_presentation'] for s in slides.slides] == [1, 2]
    assert [s.metadata['source_slide_number'] for s in slides.slides] == [1, 1]
    assert extract_presentations(state['ppt'][0], 'legacy').num_slides == 1


def test_bad_second_presentation_cleans_entire_upload(api_client):
    api, client = api_client
    auth = register(client, 'invalid-multi@example.com')
    files = [('audio', ('first.wav', audio(.1))),
             ('ppt', ('first.pptx', deck('Valid'))),
             ('ppt', ('broken.pptx', b'bad presentation'))]
    response = client.post('/api/analyze', headers=auth, files=files)
    assert response.status_code == 422 and 'broken.pptx' in response.json()['detail']
    assert not list((api.RUNTIME / 'jobs').glob('*.json'))
    assert not list((api.RUNTIME / 'uploads').glob('*'))


def test_full_preview_includes_early_passages_on_fresh_chunk_and_cached_runs(tmp_path):
    path = tmp_path / 'speech.wav'; sf.write(path, np.full(16000 * 22, .01), 16000)
    diar = tmp_path / 'diar.json'; diar.write_text('{"segments": []}')
    segments = [SimpleNamespace(start=i / 10, end=i / 10 + .05,
        text=f'passage {i}', words=None, avg_logprob=-.1) for i in range(1, 71)]
    model = SimpleNamespace(transcribe=lambda *a, **kw: (iter(segments), SimpleNamespace(language='en')))
    from src.utils.config import get
    settings = lambda key, default=None: 10 if key == 'asr.chunk_seconds' else get(key, default)
    snapshots = []
    callback = lambda passages, details: snapshots.append([s.segment_id for s in passages])
    with patch('src.audio.asr._load_whisper_model', return_value=model), patch('src.audio.asr.get', side_effect=settings):
        result, _, _ = run_asr(path, diar, 'preview', tmp_path, on_chunk=callback)
        assert len(snapshots[0]) > 50
        assert all(s[0] == 1 for s in snapshots)
        assert len(snapshots[-1]) == len(result.segments)
        snapshots.clear()
        run_asr(path, diar, 'preview', tmp_path, on_chunk=callback)
        assert snapshots == [list(range(1, len(result.segments) + 1))]
        (tmp_path / 'preview_transcript.json').unlink()
        snapshots.clear()
        run_asr(path, diar, 'preview', tmp_path, on_chunk=callback)
        assert all(s[0] == 1 for s in snapshots)
